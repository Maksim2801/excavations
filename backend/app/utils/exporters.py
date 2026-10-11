import io
from html import escape
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_RIGHT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Image as RLImage
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

_FONTS_READY = False


def _font_files() -> tuple[str, str]:
    candidates = [
        ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
        ("/usr/share/fonts/dejavu/DejaVuSans.ttf", "/usr/share/fonts/dejavu/DejaVuSans-Bold.ttf"),
        (r"C:\Windows\Fonts\arial.ttf", r"C:\Windows\Fonts\arialbd.ttf"),
    ]
    for regular, bold in candidates:
        if Path(regular).exists() and Path(bold).exists():
            return regular, bold
    raise FileNotFoundError("Не найден шрифт с поддержкой кириллицы для PDF")


def _ensure_fonts() -> None:
    global _FONTS_READY
    if _FONTS_READY:
        return
    regular, bold = _font_files()
    pdfmetrics.registerFont(TTFont("AppSans", regular))
    pdfmetrics.registerFont(TTFont("AppSans-Bold", bold))
    _FONTS_READY = True


def _thin_border() -> Border:
    return Border(
        left=Side(style="thin", color="B7A894"),
        right=Side(style="thin", color="B7A894"),
        top=Side(style="thin", color="B7A894"),
        bottom=Side(style="thin", color="B7A894"),
    )


def build_xlsx(title: str, meta_lines: list[str], headers: list[str], rows: list[list[str]]) -> bytes:
    frame = pd.DataFrame(rows, columns=headers)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Опись"
    sheet["A1"] = title
    sheet["A1"].font = Font(name="Calibri", bold=True, size=16, color="2C2416")
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(len(headers), 1))
    row_index = 3
    for line in meta_lines:
        sheet.cell(row_index, 1, line)
        row_index += 1
    sheet.cell(row_index, 1, "Единицы: координаты - м (0,001), размеры - мм (0,1).")
    header_row = row_index + 2
    header_fill = PatternFill("solid", fgColor="E7D7C1")
    thin = _thin_border()
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(header_row, column, header)
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.border = thin
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for offset, values in enumerate(frame.itertuples(index=False), start=1):
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(header_row + offset, column, "" if value is None else str(value))
            cell.border = thin
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for column in range(1, len(headers) + 1):
        sheet.column_dimensions[get_column_letter(column)].width = 18
    sheet.freeze_panes = f"A{header_row + 1}"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.oddHeader.left.text = title
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def build_formal_xlsx(
    *,
    title: str,
    approver_role: str,
    approver_name: str,
    inventory_date: str,
    requisites: list[tuple[str, str]],
    headers: list[str],
    rows: list[list[str]],
) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Опись"
    width = max(len(headers), 4)
    bold = Font(name="Times New Roman", bold=True, size=12)
    normal = Font(name="Times New Roman", size=11)
    title_font = Font(name="Times New Roman", bold=True, size=13)
    right = Alignment(horizontal="right", wrap_text=True)
    center = Alignment(horizontal="center", wrap_text=True, vertical="center")
    wrap = Alignment(wrap_text=True, vertical="top")

    sheet.merge_cells(start_row=1, start_column=max(1, width - 2), end_row=1, end_column=width)
    sheet.cell(1, max(1, width - 2), "УТВЕРЖДАЮ").font = bold
    sheet.cell(1, max(1, width - 2)).alignment = right
    sheet.merge_cells(start_row=2, start_column=max(1, width - 2), end_row=2, end_column=width)
    sheet.cell(2, max(1, width - 2), approver_role or "").font = normal
    sheet.cell(2, max(1, width - 2)).alignment = right
    sheet.merge_cells(start_row=3, start_column=max(1, width - 2), end_row=3, end_column=width)
    sheet.cell(3, max(1, width - 2), approver_name or "").font = normal
    sheet.cell(3, max(1, width - 2)).alignment = right
    sheet.merge_cells(start_row=4, start_column=max(1, width - 2), end_row=4, end_column=width)
    sheet.cell(4, max(1, width - 2), "_______________________").font = normal
    sheet.cell(4, max(1, width - 2)).alignment = right
    sheet.merge_cells(start_row=5, start_column=max(1, width - 2), end_row=5, end_column=width)
    date_line = inventory_date or "«__»_____________ ____ г."
    sheet.cell(5, max(1, width - 2), date_line).font = normal
    sheet.cell(5, max(1, width - 2)).alignment = right

    sheet.merge_cells(start_row=7, start_column=1, end_row=7, end_column=width)
    sheet.cell(7, 1, title).font = title_font
    sheet.cell(7, 1).alignment = center

    row_index = 9
    for label, value in requisites:
        sheet.cell(row_index, 1, label).font = bold
        sheet.merge_cells(start_row=row_index, start_column=1, end_row=row_index, end_column=2)
        sheet.merge_cells(start_row=row_index, start_column=3, end_row=row_index, end_column=width)
        sheet.cell(row_index, 3, value).font = normal
        sheet.cell(row_index, 1).alignment = wrap
        sheet.cell(row_index, 3).alignment = wrap
        sheet.row_dimensions[row_index].height = max(18, 14 * (1 + str(value).count("\n")))
        row_index += 1

    header_row = row_index + 1
    thin = Border(
        left=Side(style="thin", color="333333"),
        right=Side(style="thin", color="333333"),
        top=Side(style="thin", color="333333"),
        bottom=Side(style="thin", color="333333"),
    )
    header_fill = PatternFill("solid", fgColor="F3EFE6")
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(header_row, column, header)
        cell.font = Font(name="Times New Roman", bold=True, size=10)
        cell.fill = header_fill
        cell.border = thin
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    for offset, values in enumerate(rows, start=1):
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(header_row + offset, column, "" if value is None else str(value))
            cell.font = Font(name="Times New Roman", size=10)
            cell.border = thin
            cell.alignment = wrap
    widths = [6, 28, 12, 12, 12, 12, 10, 10, 14, 12, 10, 8, 14, 14]
    for column in range(1, len(headers) + 1):
        sheet.column_dimensions[get_column_letter(column)].width = widths[column - 1] if column <= len(widths) else 12
    sheet.page_setup.orientation = "portrait" if len(headers) <= 8 else "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def build_pdf(title: str, meta_lines: list[str], headers: list[str], rows: list[list[str]], generated_at: str) -> bytes:
    _ensure_fonts()
    buffer = io.BytesIO()
    document = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        leftMargin=24,
        rightMargin=24,
        topMargin=28,
        bottomMargin=28,
        title=title,
    )
    title_style = ParagraphStyle("TitleRu", fontName="AppSans-Bold", fontSize=14, leading=18, alignment=1)
    meta_style = ParagraphStyle("MetaRu", fontName="AppSans", fontSize=9, leading=12)
    cell_style = ParagraphStyle("CellRu", fontName="AppSans", fontSize=7, leading=9)
    header_style = ParagraphStyle("HeadRu", fontName="AppSans-Bold", fontSize=7, leading=9)
    story = [Paragraph(escape(title), title_style), Spacer(1, 8)]
    for line in meta_lines:
        story.append(Paragraph(escape(str(line)), meta_style))
    story.append(Paragraph("Единицы измерения: координаты - метры, размеры - миллиметры.", meta_style))
    story.append(Paragraph(f"Дата формирования: {generated_at}", meta_style))
    story.append(Spacer(1, 8))
    header = [Paragraph(column, header_style) for column in headers]
    body = [[Paragraph(escape(str(cell or "")), cell_style) for cell in row] for row in rows] or [[Paragraph("Записи не найдены", cell_style)]]
    table = Table([header] + body, repeatRows=1, colWidths=[document.width / max(1, len(headers))] * len(headers))
    table.setStyle(
        TableStyle(
            [
                ("FONTNAME", (0, 0), (-1, 0), "AppSans-Bold"),
                ("BACKGROUND", (0, 0), (-1, 0), colors.Color(0.91, 0.84, 0.76)),
                ("GRID", (0, 0), (-1, -1), 0.3, colors.Color(0.35, 0.3, 0.24)),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 3),
                ("RIGHTPADDING", (0, 0), (-1, -1), 3),
                ("TOPPADDING", (0, 0), (-1, -1), 2),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
            ]
        )
    )
    story.append(table)
    document.build(story)
    return buffer.getvalue()


def build_formal_pdf(
    *,
    title: str,
    approver_role: str,
    approver_name: str,
    inventory_date: str,
    requisites: list[tuple[str, str]],
    headers: list[str],
    rows: list[list[str]],
    generated_at: str,
) -> bytes:
    _ensure_fonts()
    buffer = io.BytesIO()
    pagesize = A4 if len(headers) <= 8 else landscape(A4)
    document = SimpleDocTemplate(
        buffer,
        pagesize=pagesize,
        leftMargin=28,
        rightMargin=28,
        topMargin=28,
        bottomMargin=28,
        title=title,
    )
    approve_style = ParagraphStyle("ApproveRu", fontName="AppSans-Bold", fontSize=10, leading=13, alignment=TA_RIGHT)
    approve_meta = ParagraphStyle("ApproveMetaRu", fontName="AppSans", fontSize=9, leading=12, alignment=TA_RIGHT)
    title_style = ParagraphStyle("FormalTitleRu", fontName="AppSans-Bold", fontSize=12, leading=15, alignment=TA_CENTER)
    label_style = ParagraphStyle("ReqLabelRu", fontName="AppSans-Bold", fontSize=8, leading=10)
    value_style = ParagraphStyle("ReqValueRu", fontName="AppSans", fontSize=8, leading=10)
    cell_style = ParagraphStyle("FormalCellRu", fontName="AppSans", fontSize=7, leading=9)
    header_style = ParagraphStyle("FormalHeadRu", fontName="AppSans-Bold", fontSize=7, leading=9)
    meta_style = ParagraphStyle("FormalMetaRu", fontName="AppSans", fontSize=8, leading=10)

    story = [
        Paragraph("УТВЕРЖДАЮ", approve_style),
        Paragraph(escape(approver_role or ""), approve_meta),
        Paragraph(escape(approver_name or ""), approve_meta),
        Paragraph("_______________________", approve_meta),
        Paragraph(escape(inventory_date or "«__»_____________ ____ г."), approve_meta),
        Spacer(1, 10),
        Paragraph(escape(title), title_style),
        Spacer(1, 10),
    ]
    for label, value in requisites:
        req = Table(
            [[Paragraph(escape(label), label_style), Paragraph(escape(value).replace("\n", "<br/>"), value_style)]],
            colWidths=[document.width * 0.38, document.width * 0.62],
        )
        req.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 4)]))
        story.append(req)
        story.append(Spacer(1, 4))
    story.append(Spacer(1, 8))
    header = [Paragraph(escape(column), header_style) for column in headers]
    body = [[Paragraph(escape(str(cell or "")), cell_style) for cell in row] for row in rows] or [[Paragraph("Записи не найдены", cell_style)]]
    table = Table([header] + body, repeatRows=1, colWidths=[document.width / max(1, len(headers))] * len(headers))
    table.setStyle(
        TableStyle(
            [
                ("FONTNAME", (0, 0), (-1, 0), "AppSans-Bold"),
                ("BACKGROUND", (0, 0), (-1, 0), colors.Color(0.95, 0.94, 0.9)),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.Color(0.2, 0.2, 0.2)),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 2),
                ("RIGHTPADDING", (0, 0), (-1, -1), 2),
                ("TOPPADDING", (0, 0), (-1, -1), 2),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
            ]
        )
    )
    story.append(table)
    story.append(Spacer(1, 8))
    story.append(Paragraph(f"Дата формирования: {escape(generated_at)}. Размеры в таблице — см.", meta_style))
    document.build(story)
    return buffer.getvalue()


def build_album_pdf(
    *,
    title: str,
    subtitle: str,
    figures: list[dict],
    generated_at: str,
    author: str,
    meta_lines: list[str] | None = None,
) -> bytes:
    """figures: [{image: bytes, heading: str, caption: str, details: str}]"""
    _ensure_fonts()
    buffer = io.BytesIO()
    document = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=40,
        rightMargin=40,
        topMargin=40,
        bottomMargin=40,
        title=title,
    )
    title_style = ParagraphStyle("AlbumTitle", fontName="AppSans-Bold", fontSize=16, leading=20, alignment=1)
    sub_style = ParagraphStyle("AlbumSub", fontName="AppSans", fontSize=11, leading=14, alignment=1)
    heading_style = ParagraphStyle("AlbumHeading", fontName="AppSans-Bold", fontSize=11, leading=14, alignment=1)
    caption_style = ParagraphStyle("AlbumCaption", fontName="AppSans", fontSize=10, leading=13, alignment=1)
    details_style = ParagraphStyle("AlbumDetails", fontName="AppSans", fontSize=9, leading=12, alignment=0)
    meta_style = ParagraphStyle("AlbumMeta", fontName="AppSans", fontSize=10, leading=13)
    toc_style = ParagraphStyle("AlbumToc", fontName="AppSans", fontSize=10, leading=14)
    story = [
        Spacer(1, 80),
        Paragraph(escape(title), title_style),
        Spacer(1, 12),
    ]
    if subtitle:
        story.append(Paragraph(escape(subtitle), sub_style))
        story.append(Spacer(1, 16))
    for line in meta_lines or []:
        if line:
            story.append(Paragraph(escape(str(line)), meta_style))
    story.extend(
        [
            Spacer(1, 18),
            Paragraph(f"Дата формирования: {escape(generated_at)}", meta_style),
            Paragraph(f"Сформировал: {escape(author)}", meta_style),
            Paragraph(f"Иллюстраций: {len(figures)}", meta_style),
            PageBreak(),
            Paragraph("Список иллюстраций", title_style),
            Spacer(1, 12),
        ]
    )
    for figure in figures:
        heading = figure.get("heading") or figure.get("caption") or ""
        story.append(Paragraph(escape(heading), toc_style))
    story.append(PageBreak())

    max_w = document.width
    max_h = document.height * 0.58
    for index, figure in enumerate(figures):
        image_bytes = figure.get("image") or b""
        heading = figure.get("heading") or ""
        caption = figure.get("caption") or ""
        details = figure.get("details") or ""
        if heading:
            story.append(Paragraph(escape(heading), heading_style))
            story.append(Spacer(1, 8))
        if image_bytes:
            img_buffer = io.BytesIO(image_bytes)
            image = RLImage(img_buffer)
            image.hAlign = "CENTER"
            iw, ih = float(image.imageWidth), float(image.imageHeight)
            scale = min(max_w / iw, max_h / ih, 1.0) if iw and ih else 1.0
            image.drawWidth = iw * scale
            image.drawHeight = ih * scale
            story.append(image)
            story.append(Spacer(1, 10))
        if caption:
            story.append(Paragraph(escape(caption), caption_style))
            story.append(Spacer(1, 6))
        if details:
            for line in str(details).split("\n"):
                if line.strip():
                    story.append(Paragraph(escape(line.strip()), details_style))
        if index < len(figures) - 1:
            story.append(PageBreak())
    if not figures:
        story.append(Paragraph("Нет иллюстраций для выбранной выборки.", caption_style))
    document.build(story)
    return buffer.getvalue()
