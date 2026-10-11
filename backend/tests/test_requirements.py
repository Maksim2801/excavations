"""Regression checks without PostGIS; spatial integration stays in test_api.py."""
from datetime import datetime, timezone
from types import SimpleNamespace

import pytest
from sqlalchemy import Column, MetaData, Table, Text, create_engine, event
from sqlalchemy.orm import Session
from geoalchemy2 import Geometry

from app.database import Base
from app.models import User, Excavation, ExcavationMember, Finding, Layer, Photo
from app.repositories.findings import apply_filters, base_query
from app.schemas.io import ReportIn
from app.services.report_service import preview_report, export_report
from app.services.finding_service import bind_gnss
from app.errors import AppError


@pytest.fixture
def local_db():
    # Plain SQLite columns only for non-spatial behaviour. Never substitutes for PostGIS tests.
    engine = create_engine("sqlite://")
    @event.listens_for(engine, "connect")
    def functions(connection, _):
        connection.create_function("ST_AsEWKB", 1, lambda value: value)
        connection.create_function("AsEWKB", 1, lambda value: value)
        connection.create_function("GeomFromEWKT", 1, lambda value: value)
    metadata = MetaData()
    for table in Base.metadata.tables.values():
        Table(table.name, metadata, *[
            Column(column.name, Text() if isinstance(column.type, Geometry) else column.type,
                   primary_key=column.primary_key, nullable=column.nullable)
            for column in table.columns
        ])
    metadata.create_all(engine)
    with Session(engine) as db:
        user = User(username="field", full_name="Field", role="field_archaeologist", hashed_password="unused")
        db.add(user)
        first = Excavation(code="A", year=2026, center_north=0, center_east=0, center_elev=0)
        second = Excavation(code="B", year=2026, center_north=0, center_east=0, center_elev=0)
        db.add_all([first, second]); db.flush()
        db.add(ExcavationMember(user_id=user.id, excavation_id=first.id))
        for eid, code, photo, day in [(first.id, "visible", False, 1), (second.id, "hidden", True, 2)]:
            db.add(Finding(excavation_id=eid, code=code, number=1, material="керамика", description="Осколок", photo_present=photo,
                           created_at=datetime(2026, 10, day, 12, tzinfo=timezone.utc)))
            db.add(Layer(excavation_id=eid, number=1, top_elev=1, bottom_elev=0))
            db.add(Photo(excavation_id=eid, file_path="/test", original_name=code, photo_type="excavation"))
        db.commit()
    with Session(engine) as db:
        yield db


def test_field_scope_covers_lists_and_direct_lookups(local_db):
    db = local_db
    db.info["actor"] = {"id": db.query(User).one().id, "role": "field_archaeologist"}
    assert [x.code for x in db.query(Excavation).all()] == ["A"]
    assert [x.code for x in base_query(db).all()] == ["visible"]
    assert db.get(Finding, 2) is None
    assert db.get(Excavation, 2) is None
    assert db.get(Layer, 2) is None
    assert db.get(Photo, 2) is None
    assert db.query(Finding).count() == 1
    db.commit()
    assert db.query(Finding).count() == 1


def test_no_assignments_means_no_field_access(local_db):
    local_db.query(ExcavationMember).delete(); local_db.commit()
    local_db.info["actor"] = {"id": local_db.query(User).one().id, "role": "field_archaeologist"}
    assert local_db.query(Excavation).count() == 0
    assert local_db.query(Finding).count() == 0


def test_inventory_and_report_filters_select_same_records(local_db):
    filters = dict(q="Осколок", photo_present=False, date_from="2026-10-01", date_to="2026-10-01", ids=[1, 2])
    inventory = apply_filters(base_query(local_db), filters).all()
    request = ReportIn(**filters, columns=["code", "description"])
    actor = SimpleNamespace(id=1, full_name="Tester")
    preview = preview_report(local_db, request, actor)
    assert preview["total"] == len(inventory) == 1
    assert preview["rows"][0][0] == inventory[0].code == "visible"
    assert apply_filters(base_query(local_db), {"ids": []}).count() == 0


def test_report_inclusive_end_date_and_reversed_range(local_db):
    actor = SimpleNamespace(id=1, full_name="Tester")
    report = preview_report(local_db, ReportIn(date_to="2026-10-02"), actor)
    assert report["total"] == 2
    with pytest.raises(AppError):
        preview_report(local_db, ReportIn(date_from="2026-10-03", date_to="2026-10-01"), actor)


def test_report_meta_and_pdf_escape_markup(local_db):
    actor = SimpleNamespace(id=1, full_name="Tester")
    request = ReportIn(format="pdf", title="A < B & C", collection_name="Коллекция", monument="Памятник", source="Поле", dating="X век", culture="Культура", columns=["code"])
    preview = preview_report(local_db, request, actor)
    assert "Коллекция" in preview["html"] and "Культура" in preview["html"]
    content, _, mime, _ = export_report(local_db, request, actor)
    assert content.startswith(b"%PDF") and mime == "application/pdf"


def test_formal_field_and_collection_layouts(local_db):
    from openpyxl import load_workbook
    from io import BytesIO

    actor = SimpleNamespace(id=1, full_name="Tester")
    finding = local_db.get(Finding, 1)
    finding.length_mm = 35
    finding.width_mm = 20
    finding.thickness_mm = 6
    finding.layer_id = 1
    finding.square = "А1"
    local_db.commit()

    field = ReportIn(
        layout="field",
        format="xlsx",
        inventory_no="A-1-2",
        inventory_date="11.10.2026",
        approver_role="директор",
        approver_name="Иванов И.И.",
        collection_name="Тестовая коллекция",
        monument="Памятник-1",
        excavation_author="Петров",
        place_time="2026 г.",
        source="Открытый лист",
        dating="XIV в.",
        culture="золотоордынская",
    )
    preview = preview_report(local_db, field, actor)
    assert preview["layout"] == "field"
    assert preview["columns"][0]["label"] == "№ п/п"
    assert "УТВЕРЖДАЮ" in preview["html"]
    assert "ПОЛЕВАЯ ОПИСЬ" in preview["title"]
    assert preview["rows"][0][1] == "Осколок"
    content, filename, mime, _ = export_report(local_db, field, actor)
    assert filename.startswith("polevaya_") and mime.endswith("sheet")
    workbook = load_workbook(BytesIO(content))
    sheet = workbook.active
    values = [cell.value for row in sheet.iter_rows(max_row=8, values_only=False) for cell in row if cell.value]
    assert "УТВЕРЖДАЮ" in values
    assert any(isinstance(value, str) and "ПОЛЕВАЯ ОПИСЬ" in value for value in values)

    collection = ReportIn(layout="collection", format="pdf", collection_name="Тестовая коллекция", monument="Памятник-1")
    preview = preview_report(local_db, collection, actor)
    assert preview["layout"] == "collection"
    assert any(column["label"] == "Полевой шифр" for column in preview["columns"])
    assert preview["rows"][0][2] == "3,5×2×0,6"
    assert "A/1/1" in preview["rows"][0]
    pdf, name, pdf_mime, _ = export_report(local_db, collection, actor)
    assert pdf.startswith(b"%PDF") and name.startswith("kollekcionnaya_") and pdf_mime == "application/pdf"

    working = preview_report(local_db, ReportIn(columns=["code", "description"]), actor)
    assert working["layout"] == "working"
    assert working["columns"][0]["key"] == "code"


def test_album_pdf_builders_and_export(local_db, monkeypatch, tmp_path):
    from PIL import Image
    from app.utils.exporters import build_album_pdf
    from app.utils.plan_render import render_plan_png
    from app.utils.files import upload_root
    import app.services.report_service as report_service

    plan = {
        "excavation_code": "A",
        "year": 2026,
        "bounds": {"min_x": 0, "max_x": 10, "min_y": 0, "max_y": 10},
        "outline": [{"x": 1, "y": 1}, {"x": 9, "y": 1}, {"x": 9, "y": 9}, {"x": 1, "y": 9}],
        "points": [{"x": 5, "y": 5, "color": "#2f6b4f"}],
        "gnss_points": [{"x": 1, "y": 1, "point_type": "corner"}],
        "legend": [{"key": "керамика", "color": "#2f6b4f"}],
    }
    png = render_plan_png(plan, width=400, height=300)
    assert png.startswith(b"\x89PNG")
    album = build_album_pdf(
        title="Альбом",
        subtitle="Тест",
        meta_lines=["Памятник: Тест"],
        figures=[{"image": png, "heading": "Рис. 1. План", "caption": "ОАН «Тест»", "details": "Точек: 1"}],
        generated_at="11.10.2026",
        author="Tester",
    )
    assert album.startswith(b"%PDF")

    upload = upload_root()
    image_name = "album_test.jpg"
    Image.new("RGB", (80, 60), (180, 120, 80)).save(upload / image_name, format="JPEG")
    photo = local_db.query(Photo).filter(Photo.excavation_id == 1).first()
    photo.finding_id = 1
    photo.file_path = image_name
    photo.photo_type = "finding"
    local_db.commit()

    monkeypatch.setattr(report_service, "build_plan", lambda *args, **kwargs: plan)
    actor = SimpleNamespace(id=1, full_name="Tester")
    with pytest.raises(AppError):
        export_report(local_db, ReportIn(layout="album", format="pdf"), actor)
    content, filename, mime, report_id = export_report(
        local_db,
        ReportIn(layout="album", format="pdf", excavation_id=1),
        actor,
    )
    assert content.startswith(b"%PDF") and filename.startswith("album_") and mime == "application/pdf"
    assert report_id > 0
    preview = preview_report(
        local_db,
        ReportIn(layout="album", excavation_id=1, monument="Памятник-1", place_time="Уфа, 2026"),
        actor,
    )
    assert preview["layout"] == "album"
    assert preview["total"] >= 2
    assert "Рис. 1." in preview["rows"][0][1]
    assert "Памятник" in preview["html"] and "Список иллюстраций" in preview["html"]


def test_automatic_number_and_full_description(local_db):
    from app.schemas.io import FindingIn
    from app.services.finding_service import create_finding
    actor = local_db.query(User).one()
    first = create_finding(local_db, FindingIn(excavation_id=1, full_desc="Камеральное описание"), actor)
    second = create_finding(local_db, FindingIn(excavation_id=1), actor)
    assert first["number"] == 2 and second["number"] == 3
    assert first["code"] == "A-2026-2"
    assert first["full_desc"] == "Камеральное описание"


def test_gnss_provenance_and_foreign_excavation(local_db):
    from app.models import GNSSPoint
    local_db.add(GNSSPoint(excavation_id=1, north=10, east=20, elev=30, point_type="control"))
    local_db.commit()
    own = local_db.get(Finding, 1)
    bind_gnss(local_db, own, 1)
    assert (own.north, own.east, own.elev, own.gnss_point_id) == (10, 20, 30, 1)
    with pytest.raises(AppError):
        bind_gnss(local_db, local_db.get(Finding, 2), 1)
    bind_gnss(local_db, own, None)
    assert own.gnss_point_id is None


def test_idempotency_rejects_changed_payload(local_db):
    from app.services.idempotency import remember, replay
    body = {"excavation_id": 1, "code": "A"}
    remember(local_db, 1, "POST /api/findings", "same", 200, {"id": 1, "excavation_id": 1}, body)
    assert replay(local_db, 1, "POST /api/findings", "same", body).status_code == 200
    with pytest.raises(AppError):
        replay(local_db, 1, "POST /api/findings", "same", {**body, "code": "B"})


def test_photo_sort_is_category_then_upload_order(local_db):
    from app.services.photo_service import list_photos
    local_db.add(Photo(excavation_id=1, file_path="/test", original_name="detail", photo_type="detail"))
    local_db.commit()
    rows = list_photos(local_db, excavation_id=1)
    assert [row["photo_type"] for row in rows] == ["detail", "excavation"]
